from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.services.output_templates import OutputTemplate


class DocumentRenderError(RuntimeError):
    pass


SUBJECT_LABELS = {
    "HISTORY": "История",
    "SOCIAL_STUDIES": "Обществознание",
    "INDIVIDUAL_PROJECT": "Индивидуальный проект",
}


def _set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


class DocxArtifactRenderer:
    def __init__(self, template: OutputTemplate):
        self.template = template

    def render(self, artifact: dict, audience: str, output_path: Path) -> Path:
        doc = Document()
        self._setup_document(doc, artifact, audience)
        kind = artifact.get("artifact_type")
        dispatch = {
            "lesson": self._render_lesson,
            "assessment": self._render_assessment,
            "worksheet": self._render_worksheet,
            "homework": self._render_homework,
            "idea_enhancement": self._render_idea,
        }
        if kind not in dispatch:
            raise DocumentRenderError(f"Unsupported artifact_type: {kind}")
        dispatch[kind](doc, artifact, audience)
        self._render_sources(doc, artifact, audience)
        self._render_footer(doc, artifact, audience)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)
        return output_path

    def _setup_document(self, doc: Document, artifact: dict, audience: str):
        section = doc.sections[0]
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        m = Cm(self.template.margin_cm)
        section.top_margin = m
        section.bottom_margin = m
        section.left_margin = m
        section.right_margin = m

        styles = doc.styles
        normal = styles["Normal"]
        normal.font.name = self.template.font_name
        normal._element.rPr.rFonts.set(qn("w:ascii"), self.template.font_name)
        normal._element.rPr.rFonts.set(qn("w:hAnsi"), self.template.font_name)
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), self.template.font_name)
        normal.font.size = Pt(self.template.body_size_pt)
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.05

        for name, size in (
            ("Title", self.template.title_size_pt),
            ("Heading 1", self.template.heading_size_pt),
            ("Heading 2", self.template.heading_size_pt - 1),
        ):
            st = styles[name]
            st.font.name = self.template.font_name
            st._element.rPr.rFonts.set(qn("w:ascii"), self.template.font_name)
            st._element.rPr.rFonts.set(qn("w:hAnsi"), self.template.font_name)
            st._element.rPr.rFonts.set(qn("w:eastAsia"), self.template.font_name)
            st.font.size = Pt(size)
            st.font.color.rgb = RGBColor(31, 54, 82)

        # Header identity strip.
        header = section.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("УЧИТЕЛЬСКАЯ ВЕРСИЯ" if audience == "teacher" else "ВЕРСИЯ УЧЕНИКА")
        r.bold = True
        r.font.name = self.template.font_name
        r.font.size = Pt(self.template.small_size_pt)
        r.font.color.rgb = RGBColor(89, 89, 89)

    def _title(self, doc: Document, artifact: dict, audience: str):
        p = doc.add_paragraph(style="Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(artifact.get("title") or "Учебный материал")
        run.bold = True

        grade = artifact.get("grade")
        subject = SUBJECT_LABELS.get(artifact.get("subject"), artifact.get("subject") or "")
        meta = []
        if grade:
            meta.append(f"{grade} класс")
        if subject:
            meta.append(subject)
        if meta:
            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p2.add_run(" · ".join(meta))
            r.font.size = Pt(self.template.body_size_pt)
            r.bold = True
            r.font.color.rgb = RGBColor(68, 68, 68)

        if audience == "student":
            table = doc.add_table(rows=1, cols=3)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            labels = ["ФИО: __________________", "Класс: _____", "Дата: _________"]
            widths = [8.5, 4.0, 4.5]
            for idx, (cell, label, width) in enumerate(zip(table.rows[0].cells, labels, widths)):
                table.columns[idx].width = Cm(width)
                cell.width = Cm(width)
                p = cell.paragraphs[0]
                run = p.add_run(label)
                run.font.size = Pt(self.template.body_size_pt - 1)
                _set_cell_margins(cell, 40, 40, 40, 40)
            doc.add_paragraph()

    def _heading(self, doc, text, level=1):
        doc.add_paragraph(text, style=f"Heading {level}")

    def _bullet_list(self, doc, items):
        for item in items or []:
            doc.add_paragraph(str(item), style="List Bullet")

    def _numbered_list(self, doc, items):
        for item in items or []:
            doc.add_paragraph(str(item), style="List Number")

    def _info_box(self, doc, title, lines, fill="EEF3F8"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        _set_cell_shading(cell, fill)
        _set_cell_margins(cell, 100, 140, 100, 140)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        for line in lines or []:
            p2 = cell.add_paragraph(str(line))
            p2.paragraph_format.space_after = Pt(2)

    def _render_lesson(self, doc, a, audience):
        self._title(doc, a, audience)
        self._info_box(doc, "Паспорт урока", [
            f"Продолжительность: {a.get('duration_minutes', 45)} мин.",
            f"Цель: {a.get('goal', '')}",
        ])
        if a.get("essential_question"):
            self._heading(doc, "Проблемный вопрос")
            doc.add_paragraph(a["essential_question"])
        self._heading(doc, "Планируемые результаты")
        self._bullet_list(doc, a.get("measurable_outcomes", []))
        self._heading(doc, "Ключевые понятия")
        doc.add_paragraph(", ".join(a.get("key_concepts", [])) or "-")
        self._heading(doc, "Ход урока")
        for step in a.get("timeline", []):
            table = doc.add_table(rows=1, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            table.cell(0,0).width = Cm(3.0)
            table.cell(0,1).width = Cm(14.0)
            _set_cell_shading(table.cell(0,0), "D9E5F1")
            _set_cell_shading(table.cell(0,1), "F7F9FB")
            table.cell(0,0).paragraphs[0].add_run(f"{step.get('minutes')} мин.\n{step.get('stage','')}").bold = True
            body = table.cell(0,1)
            body.paragraphs[0].add_run("Учитель: ").bold = True
            body.paragraphs[0].add_run(step.get("teacher_action", ""))
            p = body.add_paragraph()
            p.add_run("Ученики: ").bold = True
            p.add_run(step.get("student_action", ""))
            if step.get("evidence_of_learning"):
                p2 = body.add_paragraph()
                p2.add_run("Свидетельство понимания: ").bold = True
                p2.add_run(step["evidence_of_learning"])
            for cell in table.rows[0].cells:
                _set_cell_margins(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            doc.add_paragraph()
        if a.get("notebook"):
            self._heading(doc, "В тетрадь")
            self._bullet_list(doc, a["notebook"])
        if a.get("differentiation"):
            self._heading(doc, "Дифференциация")
            self._bullet_list(doc, a["differentiation"])
        if a.get("formative_check"):
            self._heading(doc, "Проверка понимания")
            self._bullet_list(doc, a["formative_check"])
        if a.get("reflection"):
            self._heading(doc, "Рефлексия")
            self._bullet_list(doc, a["reflection"])
        if a.get("homework"):
            self._heading(doc, "Домашнее задание")
            doc.add_paragraph(a["homework"])
        if audience == "teacher" and a.get("cut_first_if_short_on_time"):
            self._heading(doc, "Если не хватает времени")
            self._bullet_list(doc, a["cut_first_if_short_on_time"])

    def _render_assessment(self, doc, a, audience):
        self._title(doc, a, audience)
        self._info_box(doc, "Инструкция", [
            a.get("instructions", ""),
            f"Время: {a.get('time_minutes','-')} мин.",
            f"Максимум: {a.get('max_points','-')} баллов.",
        ])
        doc.add_paragraph()
        for task in a.get("tasks", []):
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            r = p.add_run(f"{task.get('number')}. {task.get('instruction','')}")
            r.bold = True
            if task.get("points") is not None:
                r2 = p.add_run(f"  [{task.get('points')} б.]")
                r2.font.size = Pt(self.template.small_size_pt)
                r2.font.color.rgb = RGBColor(89,89,89)
            if task.get("student_data"):
                doc.add_paragraph(task["student_data"])
            if audience == "student":
                for _ in range(2 if not self.template.compact else 1):
                    doc.add_paragraph("________________________________________________________________________________")
            else:
                if task.get("answer") is not None:
                    self._info_box(doc, "Ответ", [task.get("answer", "")], fill="EAF4EA")
                if task.get("criteria") is not None:
                    self._info_box(doc, "Критерии", [task.get("criteria", "")], fill="FFF4DB")
            doc.add_paragraph()
        if audience == "teacher" and a.get("grading_scale"):
            self._heading(doc, "Шкала оценивания")
            table = doc.add_table(rows=1, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            hdr[0].text = "Отметка"
            hdr[1].text = "Диапазон"
            _set_repeat_table_header(table.rows[0])
            for c in hdr:
                _set_cell_shading(c, "D9E5F1")
            for grade, rule in a["grading_scale"].items():
                row = table.add_row().cells
                row[0].text = str(grade)
                row[1].text = str(rule)
        if audience == "teacher" and a.get("common_errors_to_watch"):
            self._heading(doc, "Типичные ошибки")
            self._bullet_list(doc, a["common_errors_to_watch"])

    def _render_worksheet(self, doc, a, audience):
        self._title(doc, a, audience)
        self._info_box(doc, "Задача рабочего листа", [a.get("purpose", "")])
        if a.get("support_box"):
            self._info_box(doc, "Опора", a["support_box"], fill="F2F2F2")
        for section in a.get("sections", []):
            self._heading(doc, section.get("title", "Раздел"))
            if section.get("instruction"):
                p = doc.add_paragraph(section["instruction"])
                p.runs[0].italic = True
            for task in section.get("tasks", []):
                p = doc.add_paragraph()
                p.add_run(f"{task.get('number')}. ").bold = True
                p.add_run(task.get("instruction", ""))
                if audience == "student":
                    doc.add_paragraph("________________________________________________________________________________")
                    doc.add_paragraph("________________________________________________________________________________")
                elif task.get("answer"):
                    self._info_box(doc, "Ключ", [task["answer"]], fill="EAF4EA")
        if a.get("reflection"):
            self._heading(doc, "Рефлексия")
            self._numbered_list(doc, a["reflection"])
        if audience == "teacher" and a.get("teacher_key"):
            self._heading(doc, "Учителю")
            self._bullet_list(doc, a["teacher_key"])

    def _render_homework(self, doc, a, audience):
        self._title(doc, a, audience)
        self._info_box(doc, "Параметры", [
            f"Цель: {a.get('purpose','')}",
            f"Ориентировочное время: {a.get('estimated_minutes','-')} мин.",
        ])
        self._heading(doc, "Задание")
        doc.add_paragraph(a.get("student_instruction", ""))
        self._heading(doc, "Критерии успеха")
        for item in a.get("success_criteria", []):
            doc.add_paragraph(f"☐ {item}")
        if a.get("optional_challenge"):
            self._heading(doc, "По желанию")
            doc.add_paragraph(a["optional_challenge"])
        if audience == "teacher" and a.get("teacher_note"):
            self._info_box(doc, "Учителю", [a["teacher_note"]], fill="FFF4DB")

    def _render_idea(self, doc, a, audience):
        self._title(doc, {**a, "title": "Развитие педагогической идеи"}, audience)
        self._heading(doc, "Исходная идея")
        doc.add_paragraph(a.get("original_idea", ""))
        self._heading(doc, "Педагогическая ценность")
        self._bullet_list(doc, a.get("pedagogical_value", []))
        self._heading(doc, "Что стоит усилить")
        self._bullet_list(doc, a.get("weaknesses_or_risks", []))
        self._info_box(doc, "Рекомендуемый результат", [a.get("recommended_result", "")])
        for opt in a.get("options", []):
            label = opt.get("label", "option").upper()
            self._heading(doc, f"Вариант: {label}")
            doc.add_paragraph(opt.get("description", ""))
            self._bullet_list(doc, opt.get("implementation", []))
            if opt.get("risks"):
                p = doc.add_paragraph()
                p.add_run("Риски: ").bold = True
                p.add_run("; ".join(opt["risks"]))
        self._info_box(doc, "Рекомендация", [
            f"Выбрать вариант: {a.get('recommended_option','')}",
            a.get("preserve_core", ""),
        ], fill="EAF4EA")

    def _render_sources(self, doc, a, audience):
        refs = a.get("source_references") or []
        if audience != "teacher" or not refs:
            return
        self._heading(doc, "Источники, использованные в материале")
        for idx, ref in enumerate(refs, 1):
            parts = [ref.get("title", "Источник")]
            if ref.get("section"):
                parts.append(ref["section"])
            if ref.get("page_print_start") is not None:
                p1 = ref["page_print_start"]
                p2 = ref.get("page_print_end")
                parts.append(f"с. {p1}" if not p2 or p2 == p1 else f"с. {p1}-{p2}")
            elif ref.get("page_pdf_start") is not None:
                p1 = ref["page_pdf_start"]
                p2 = ref.get("page_pdf_end")
                parts.append(f"PDF-стр. {p1}" if not p2 or p2 == p1 else f"PDF-стр. {p1}-{p2}")
            doc.add_paragraph(f"{idx}. " + ", ".join(str(x) for x in parts))

    def _render_footer(self, doc, artifact, audience):
        footer = doc.sections[0].footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(
            f"{self.template.name} v{self.template.version} · "
            + ("версия учителя" if audience == "teacher" else "версия ученика")
        )
        run.font.size = Pt(self.template.small_size_pt)
        run.font.color.rgb = RGBColor(117,117,117)
